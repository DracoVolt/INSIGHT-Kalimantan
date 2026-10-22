"""Build the Word version of the INSIGHT manuscript from its Markdown source.

Usage:
    python docs/paper/build_docx.py

Requires ``python-docx`` (listed in requirements-dev.txt). The converter only
understands the subset of Markdown used in ``INSIGHT-naskah-revisi.md``:
headings, paragraphs, bullet/numbered lists, pipe tables, bold, italic and
inline code.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "INSIGHT-naskah-revisi.md"
TARGET = HERE / "INSIGHT-naskah-revisi.docx"

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def add_inline(paragraph, text: str) -> None:
    """Add text to a paragraph, honouring **bold**, *italic* and `code`."""
    chunks = INLINE.split(text)
    for chunk in chunks:
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            run = paragraph.add_run(chunk[1:-1])
            run.italic = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(chunk)


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    cells = split_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c) for c in cells)


def style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    fmt = normal.paragraph_format
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.15
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)


def build() -> Path:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line or line == "---":
            i += 1
            continue

        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not is_separator(lines[i]):
                    rows.append(split_row(lines[i]))
                i += 1
            if not rows:
                continue
            n_cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=n_cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row in enumerate(rows):
                for c_idx in range(n_cols):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    text = row[c_idx] if c_idx < len(row) else ""
                    add_inline(paragraph, text)
                    paragraph.paragraph_format.space_after = Pt(2)
                    if r_idx == 0:
                        for run in paragraph.runs:
                            run.bold = True
            doc.add_paragraph()
            continue

        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line[2:].strip())
            run.bold = True
            run.font.size = Pt(12)
            i += 1
            continue
        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:].strip())
            i += 1
            continue
        if re.match(r"^\d+\.\s", line):
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", line))
            i += 1
            continue

        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(paragraph, line)
        i += 1

    doc.save(TARGET)
    return TARGET


if __name__ == "__main__":
    path = build()
    print(f"wrote {path}")
